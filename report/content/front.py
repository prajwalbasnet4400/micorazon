"""Cover, title page, certificate pages, acknowledgement, abstract, indexes, abbreviations."""

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .common import ADDRESS, COLLEGE, META, ORIG, SUBMIT_DATE, SUPERVISOR, TEAM, UNIVERSITY, pct


def _title_page(R):
    p = R.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ORIG / "college_logo.png"), width=Inches(1.15))
    R.para(COLLEGE.upper(), align="center", size=14, bold=True, space_after=0)
    R.para(ADDRESS, align="center", space_after=0)
    R.para(f"(Affiliated to {UNIVERSITY})", align="center", space_after=24)
    R.para("A Final Year Project Report on", align="center", size=14, space_after=4)
    R.para("“HEART FAILURE PREDICTION”", align="center", size=20, bold=True, space_after=4)
    R.para("(Project Code: BIT453CO)", align="center", space_after=24)
    R.para(
        "Submitted in partial fulfillment of the requirements for the degree of\n"
        "Bachelor in Information Technology (BIT)",
        align="center", space_after=24,
    )
    R.para("Submitted by:", align="center", bold=True, space_after=2)
    for name, roll in TEAM:
        R.para(f"{name} ({roll})", align="center", space_after=0)
    R.para("", space_after=6)
    R.para("Under the Supervision of:", align="center", bold=True, space_after=2)
    R.para(SUPERVISOR, align="center", space_after=0)
    R.para("", space_after=10)
    R.para("Submitted to:", align="center", bold=True, space_after=2)
    R.para("Department of Science and Technology", align="center", space_after=0)
    R.para(COLLEGE, align="center", space_after=18)
    R.para(SUBMIT_DATE, align="center", bold=True)


def _signature_table(R, cells):
    t = R.doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(cells), 2):
        row = t.add_row().cells
        for j, (role, name) in enumerate(cells[i:i + 2]):
            c = row[j]
            c.width = Inches(3.0)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(36)
            p.add_run("…………………………………")
            for line, bold in ((name, True), (role, False)):
                q = c.add_paragraph()
                q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                q.paragraph_format.space_after = Pt(0)
                r = q.add_run(line)
                r.bold = bold
    R._set_widths(t, [3.0, 3.0])


def front_matter(R, cite):
    team_names = ", ".join(n for n, _ in TEAM[:-1]) + f" and {TEAM[-1][0]}"
    team_rolls = ", ".join(f"{n} ({r})" for n, r in TEAM[:-1]) + f" and {TEAM[-1][0]} ({TEAM[-1][1]})"

    # Title page (unnumbered).
    R.configure_first_section(None)
    _title_page(R)

    # Certificate pages onward: roman numerals starting at i.
    R.new_section("lowerRoman", 1)
    R.heading("SUPERVISOR RECOMMENDATION", 1)
    R.para(
        f"I hereby recommend that this project report prepared under my supervision by {team_rolls}, "
        "entitled “Heart Failure Prediction”, in partial fulfillment of the requirements for the degree "
        f"of Bachelor in Information Technology (BIT) of {UNIVERSITY}, be processed for evaluation."
    )
    R.para(
        "To the best of my knowledge, the work presented in this report is the original work of the students, "
        "and the system described has been designed, implemented and tested by them."
    )
    R.para("")
    R.para("", space_after=30)
    R.para("…………………………………", align="left", space_after=0)
    R.para(f"**{SUPERVISOR}**", align="left", space_after=0)
    R.para("Project Supervisor", align="left", space_after=0)
    R.para("Department of Science and Technology", align="left", space_after=0)
    R.para(COLLEGE, align="left", space_after=0)
    R.para("Date: ……………………", align="left")

    R.heading("LETTER OF APPROVAL", 1, page_break=True)
    R.para(
        f"This is to certify that this project report prepared by {team_rolls}, entitled "
        "“Heart Failure Prediction”, in partial fulfillment of the requirements for the degree of "
        "Bachelor in Information Technology (BIT), has been evaluated. In our opinion, it is satisfactory "
        "in scope and quality as a project for the required degree."
    )
    R.para("Evaluation Committee", align="center", bold=True, space_after=0)
    _signature_table(R, [
        ("Supervisor", SUPERVISOR),
        ("Head of Department", "………………………"),
        ("Internal Examiner", "………………………"),
        ("External Examiner", "………………………"),
    ])

    R.heading("ACKNOWLEDGEMENT", 1, page_break=True)
    R.para(
        f"We would like to express our sincere gratitude to our supervisor, {SUPERVISOR}, for the valuable "
        "guidance, constructive feedback and continuous support throughout this project. The suggestions received during "
        "the proposal, mid-term and final phases helped us shape both the machine learning model and the web "
        "application."
    )
    R.para(
        f"We are thankful to the Department of Science and Technology, {COLLEGE}, for providing the opportunity, "
        "resources and environment needed to carry out this project. We also thank the faculty members who "
        "reviewed our work and offered suggestions during the presentations."
    )
    R.para(
        "We acknowledge the creators of the Heart Failure Clinical Records dataset and the maintainers of the UCI "
        "Machine Learning Repository, whose openly available data made this work possible."
    )
    R.para(
        "Lastly, we would like to thank our friends and family for their continuous motivation and support "
        "during the project journey."
    )
    R.para("")
    for name, roll in TEAM:
        R.para(f"{name} ({roll})", align="right", space_after=0)

    m = META["metrics"]
    rf = META["comparison"]["random_forest"]["test"]
    R.heading("ABSTRACT", 1, page_break=True)
    R.para(
        "Heart failure is a leading cause of hospitalization and death worldwide, and in low-resource settings "
        "such as Nepal, early identification of high-risk patients is limited by manual, experience-driven "
        "assessment. This project presents the Heart Failure Prediction (HFP) system, a web application that "
        "estimates a patient's risk of a heart-failure death event from twelve routinely collected clinical "
        "attributes, including age, ejection fraction, serum creatinine, serum sodium and comorbidities."
    )
    R.para(
        "The model was developed on the UCI Heart Failure Clinical Records dataset, expanded from 299 real "
        "patients to 2000 records by controlled jittering, with a patient-group-aware train/test split to "
        "prevent information leakage. Six classifiers were compared: Logistic Regression, Support Vector "
        "Machine, Decision Tree, Random Forest, XGBoost and Gradient Boosting. On 589 held-out records, the "
        f"deployed XGBoost model achieved an accuracy of {pct(m['accuracy'])}, precision of {pct(m['precision'])}, "
        f"recall of {pct(m['recall'])}, F1-score of {m['f1']:.3f} and the highest ROC-AUC of {m['roc_auc']:.3f}. "
        f"Random Forest obtained a slightly higher F1-score ({rf['f1']:.3f}); XGBoost was retained for its equal "
        "discrimination ability and its native, exact SHAP explanations."
    )
    R.para(
        "The system is implemented with a FastAPI backend, a React frontend and an SQLite database. Users can "
        "register, submit clinical data, and receive a risk status, probability, risk level, the five factors "
        "that most influenced their prediction and rule-based health suggestions, together with a downloadable "
        "PDF report and a history of previous predictions. The system was verified with 19 automated unit and "
        "integration tests and 14 browser-based system test cases."
    )
    R.para(
        "**Keywords:** heart failure, machine learning, XGBoost, explainable AI, SHAP, clinical decision support"
    )

    R.heading("TABLE OF CONTENTS", 1, page_break=True)
    R.toc()

    R.heading("LIST OF ABBREVIATIONS", 1, page_break=True)
    abbreviations = [
        ("AI", "Artificial Intelligence"), ("API", "Application Programming Interface"),
        ("ASGI", "Asynchronous Server Gateway Interface"), ("AUC", "Area Under the Curve"),
        ("BIT", "Bachelor in Information Technology"), ("CPK", "Creatinine Phosphokinase"),
        ("CSV", "Comma-Separated Values"), ("CT", "Computed Tomography"), ("CV", "Cross-Validation"),
        ("DL", "Deep Learning"), ("EF", "Ejection Fraction"), ("EHR", "Electronic Health Record"),
        ("FN", "False Negative"), ("FP", "False Positive"), ("HF", "Heart Failure"),
        ("HFP", "Heart Failure Prediction"), ("HTTP", "Hypertext Transfer Protocol"),
        ("IEEE", "Institute of Electrical and Electronics Engineers"), ("JSON", "JavaScript Object Notation"),
        ("JWT", "JSON Web Token"), ("LIME", "Local Interpretable Model-agnostic Explanations"),
        ("ML", "Machine Learning"), ("PBKDF2", "Password-Based Key Derivation Function 2"),
        ("PDF", "Portable Document Format"), ("REST", "Representational State Transfer"),
        ("ROC", "Receiver Operating Characteristic"), ("SHAP", "SHapley Additive exPlanations"),
        ("SPA", "Single Page Application"), ("SQL", "Structured Query Language"),
        ("SVM", "Support Vector Machine"), ("TN", "True Negative"), ("TP", "True Positive"),
        ("UCI", "University of California, Irvine"), ("UI", "User Interface"),
        ("UML", "Unified Modeling Language"), ("XGBoost", "Extreme Gradient Boosting"),
    ]
    t = R.doc.add_table(rows=0, cols=2)
    for short, full in abbreviations:
        cells = t.add_row().cells
        R._cell(cells[0], short, size=12)
        R._cell(cells[1], full, size=12)
        for c in cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(4)
    R._set_widths(t, [1.6, 4.4])

    R.heading("LIST OF FIGURES", 1, page_break=True)
    R.list_of("Figure")
    R.heading("LIST OF TABLES", 1, page_break=True)
    R.list_of("Table")
