"""Helpers for building the project report in the prescribed format.

Format (Project III guideline):
  * A4, margins top/bottom/right 1", left 1.25"
  * Times New Roman 12 pt, justified, 1.5 line spacing
  * Headings bold: chapter 16 pt, section 14 pt, sub-section 12 pt
  * Figures/tables centred; figure captions centred below, table captions centred above,
    captions bold 12 pt
  * Roman page numbers from the certificate page to the lists, Arabic from Chapter 1,
    page number at the bottom centre
"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
TEXT_WIDTH = Inches(8.27 - 1.25 - 1.0)  # A4 width minus margins


def _set_font(style_or_run, size=None, bold=None, italic=None):
    font = style_or_run.font
    font.name = FONT
    font.color.rgb = BLACK
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def _field(paragraph, instr, cached=""):
    """Insert a complex field (PAGE, SEQ, TOC ...) with an optional cached result."""
    def run_with(el):
        r = paragraph.add_run()
        r._r.append(el)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_with(begin)
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run_with(instr_el)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    result = paragraph.add_run(cached)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)
    return result


class Report:
    def __init__(self, index_pages=None):
        """index_pages: page labels for TOC/figure/table entries from a previous render pass."""
        self.doc = Document()
        self.figure_no = 0
        self.table_no = 0
        # Entries collected while building: ("toc", level, text) / ("Figure", 0, text) / ("Table", 0, text)
        self.entries: list[tuple[str, int, str]] = []
        self.index_pages = index_pages
        self._index_fields = []  # (kind, first_paragraph) placeholders filled at save()
        self.labels: dict[str, int] = {}  # "fig:name" / "tab:name" -> number, for cross-references
        self._setup_page(self.doc.sections[0])
        self._setup_styles()

    # ------------------------------------------------------------------ setup
    @staticmethod
    def _setup_page(section):
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = section.bottom_margin = section.right_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.footer_distance = Inches(0.5)

    def _setup_styles(self):
        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, 12)
        pf = normal.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)

        for name, size, align, before, after in (
            ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 0, 18),
            ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6),
            ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 4),
            ("Heading 4", 12, WD_ALIGN_PARAGRAPH.LEFT, 8, 4),
        ):
            st = styles[name]
            _set_font(st, size, bold=True, italic=False)
            st.paragraph_format.alignment = align
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            st.paragraph_format.keep_with_next = True

        caption = styles["Caption"]
        _set_font(caption, 12, bold=True, italic=False)
        caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(4)
        caption.paragraph_format.space_after = Pt(12)
        caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

        # Title-style heading that is NOT collected into the table of contents.
        plain = styles.add_style("PlainTitle", 1)
        plain.base_style = normal
        _set_font(plain, 16, bold=True)
        plain.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        plain.paragraph_format.space_after = Pt(18)

        for name, size in (("TableText", 11), ("CodeText", 9)):
            st = styles.add_style(name, 1)
            st.base_style = normal
            _set_font(st, size)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            st.paragraph_format.space_after = Pt(0)
        code = styles["CodeText"]
        code.font.name = "Courier New"
        code.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:ascii"), "Courier New")
        code.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:hAnsi"), "Courier New")

        for name, indent in (("toc 1", 0), ("toc 2", 0.3), ("toc 3", 0.6), ("table of figures", 0)):
            st = styles.add_style(name, 1)
            st.base_style = normal
            _set_font(st, 12)
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            st.paragraph_format.left_indent = Inches(indent)
            st.paragraph_format.space_after = Pt(2)
            st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            st.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        styles["toc 1"].font.bold = True

        bullet = styles["List Bullet"]
        _set_font(bullet, 12)
        bullet.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        bullet.paragraph_format.space_after = Pt(2)
        bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        numbered = styles["List Number"]
        _set_font(numbered, 12)
        numbered.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        numbered.paragraph_format.space_after = Pt(2)

    # --------------------------------------------------------------- sections
    def new_section(self, numbering: str | None, start: int = 1):
        """Start a new page section. numbering: None (no number), 'lowerRoman' or 'decimal'."""
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(section)
        self._configure_footer(section, numbering, start)
        return section

    def configure_first_section(self, numbering=None):
        self._configure_footer(self.doc.sections[0], numbering, 1)

    @staticmethod
    def _configure_footer(section, numbering, start):
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.text = ""
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        if numbering:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = _field(p, "PAGE", "1")
            _set_font(run, 12)
            pg = OxmlElement("w:pgNumType")
            pg.set(qn("w:fmt"), numbering)
            pg.set(qn("w:start"), str(start))
            section._sectPr.append(pg)

    # ---------------------------------------------------------------- content
    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def heading(self, text, level=1, page_break=False):
        h = self.doc.add_heading(level=level)
        if page_break:
            h.paragraph_format.page_break_before = True
        self._rich(h, text)
        if level <= 3:
            self.entries.append(("toc", level, re.sub(r"\*", "", text)))
        for r in h.runs:
            _set_font(r)
        return h

    def chapter(self, number, title, first=False):
        h = self.doc.add_heading(level=1)
        h.paragraph_format.page_break_before = not first
        run = h.add_run(f"CHAPTER {number}: {title.upper()}")
        _set_font(run, 16, bold=True)
        self.entries.append(("toc", 1, f"CHAPTER {number}: {title.upper()}"))
        return h

    def plain_title(self, text, page_break=False):
        p = self.doc.add_paragraph(style="PlainTitle")
        p.paragraph_format.page_break_before = page_break
        p.add_run(text)
        return p

    def para(self, text="", align=None, style=None, size=None, bold=None, space_after=None, keep_next=False):
        p = self.doc.add_paragraph(style=style)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        self._rich(p, text, size=size, bold=bold)
        return p

    @staticmethod
    def _rich(p, text, size=None, bold=None):
        """Minimal inline markup: **bold**, *italic*."""
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            if not part:
                continue
            if part.startswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("*"):
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                r = p.add_run(part)
                if bold:
                    r.bold = True
            if size:
                r.font.size = Pt(size)

    def bullets(self, items, style="List Bullet"):
        for item in items:
            p = self.doc.add_paragraph(style=style)
            self._rich(p, item)

    def numbered(self, items):
        # Manual numbering keeps lists independent (Word numbering continues across lists).
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
            self._rich(p, f"{i}.  {item}")

    # ---------------------------------------------------------------- figures
    def caption(self, kind, title):
        if kind == "Figure":
            self.figure_no += 1
            no = self.figure_no
        else:
            self.table_no += 1
            no = self.table_no
        p = self.doc.add_paragraph(style="Caption")
        r = p.add_run(f"{kind} ")
        _set_font(r, 12, bold=True)
        r = _field(p, f"SEQ {kind} \\* ARABIC", str(no))
        _set_font(r, 12, bold=True)
        r = p.add_run(f": {title}")
        _set_font(r, 12, bold=True)
        self.entries.append((kind, 0, f"{kind} {no}: {title}"))
        if kind == "Table":
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(4)
        return no

    def figure(self, path, title, width=None, label=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(6)
        p.add_run().add_picture(str(path), width=width or Inches(5.5))
        no = self.caption("Figure", title)
        if label:
            self.labels[f"fig:{label}"] = no
        return no

    def equation(self, path, height=None):
        """Insert an equation image at its natural size (rendered at 15 pt / 300 dpi -> ~13 pt)."""
        from PIL import Image

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        px_w, _ = Image.open(path).size
        p.add_run().add_picture(str(path), width=Inches(min(px_w / 300 * 0.87, TEXT_WIDTH.inches)))

    # ----------------------------------------------------------------- tables
    def table(self, title, header, rows, widths=None, font_size=11, caption=True, label=None):
        if caption:
            no = self.caption("Table", title)
            if label:
                self.labels[f"tab:{label}"] = no
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            self._cell(t.rows[0].cells[i], text, bold=True, size=font_size, shade=True)
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, text in enumerate(row):
                self._cell(cells[i], str(text), size=font_size)
        self._set_widths(t, widths)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        return t

    def kv_table(self, title, pairs, widths=(1.7, 4.3), font_size=11, label=None):
        """Two-column description table (e.g. use case descriptions)."""
        no = self.caption("Table", title)
        if label:
            self.labels[f"tab:{label}"] = no
        t = self.doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for key, value in pairs:
            cells = t.add_row().cells
            self._cell(cells[0], key, bold=True, size=font_size, shade=True)
            if isinstance(value, (list, tuple)):
                self._cell(cells[1], value[0], size=font_size)
                for extra in value[1:]:
                    p = cells[1].add_paragraph(style="TableText")
                    self._rich(p, extra, size=font_size)
            else:
                self._cell(cells[1], value, size=font_size)
        self._set_widths(t, widths)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        return t

    @staticmethod
    def _set_widths(table, widths):
        """Fixed column widths that LibreOffice honours (tblGrid + tcW + fixed layout)."""
        if not widths:
            n = len(table.columns)
            widths = [TEXT_WIDTH.inches / n] * n
        total = sum(widths)
        if total > TEXT_WIDTH.inches:  # never exceed the text block
            widths = [w * TEXT_WIDTH.inches / total for w in widths]
        table.autofit = False
        tblPr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
        for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            col.set(qn("w:w"), str(int(w * 1440)))
        for row in table.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Inches(w)

    def _cell(self, cell, text, bold=False, size=11, shade=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.style = self.doc.styles["TableText"]
        self._rich(p, text, size=size, bold=bold)
        if shade:
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "E7E6E6")
            tcpr.append(shd)

    @staticmethod
    def _repeat_header(row):
        trpr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)

    def code(self, text):
        for line in text.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph(style="CodeText")
            p.add_run(line if line else " ")

    # ---------------------------------------------------------------- indexes
    # The index fields are written at save() time, when all entries are known. Their cached
    # result holds the entries with page numbers from the previous render pass, so the document
    # is correct as delivered and can still be refreshed with "Update Index" in LibreOffice/Word.
    def toc(self):
        self._index_fields.append(("toc", self.doc.add_paragraph()))

    def list_of(self, kind):
        self._index_fields.append((kind, self.doc.add_paragraph()))

    def _fill_index(self, kind, anchor):
        instr = 'TOC \\o "1-3" \\h \\z \\u' if kind == "toc" else f'TOC \\h \\z \\c "{kind}"'
        items = [(lvl, text) for k, lvl, text in self.entries if k == kind]
        pages = (self.index_pages or {}).get(kind, [])
        lines = []
        for i, (lvl, text) in enumerate(items):
            style = f"toc {lvl}" if kind == "toc" else "table of figures"
            lines.append((style, text, pages[i] if i < len(pages) else "000"))
        if not lines:
            lines = [("Normal", "No entries", "")]

        paras = [anchor]
        for _ in lines[1:]:
            new = OxmlElement("w:p")
            paras[-1]._p.addnext(new)
            from docx.text.paragraph import Paragraph
            paras.append(Paragraph(new, anchor._parent))

        def fld(p, typ):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), typ)
            p.add_run()._r.append(el)

        for i, (p, (style, text, page)) in enumerate(zip(paras, lines)):
            p.style = self.doc.styles[style]
            if i == 0:
                fld(p, "begin")
                ins = OxmlElement("w:instrText")
                ins.set(qn("xml:space"), "preserve")
                ins.text = f" {instr} "
                p.add_run()._r.append(ins)
                fld(p, "separate")
            p.add_run(f"{text}\t{page}")
            if i == len(lines) - 1:
                fld(p, "end")

    def save(self, path):
        for kind, anchor in self._index_fields:
            self._fill_index(kind, anchor)
        self._index_fields = []
        self.doc.save(str(path))
